#!/usr/bin/env python3
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List

import requests
from openpyxl import load_workbook

EXCEL_PATH = Path(os.environ.get(
    "FINAL_POS_EXCEL",
    "config/FINAL_Production_Excel_A_WebC_Template_100Tenants.xlsx",
))
BATCH_SIZE = 100
EXPECTED_EXPORT_ROWS = 352
EXPECTED_TENANTS = 100
EXPECTED_SERVICES = 140
EXPECTED_DEVICES = 108
EXPECTED_TEMPLATES = {"BASIC", "PRO", "ENTERPRISE"}
REQUEST_TIMEOUT = 30


def _json_value(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.isoformat()
    return v


def _rows_as_dicts(ws) -> List[Dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    result: List[Dict[str, Any]] = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        result.append({headers[i]: _json_value(row[i]) for i in range(len(headers))})
    return result


def _parse_json_object(raw: Any, *, key: str) -> Dict[str, Any]:
    if raw in (None, ""):
        return {}
    if isinstance(raw, dict):
        return dict(raw)
    try:
        value = json.loads(str(raw))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid JSON for {key}: {exc}") from exc
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object for {key}, got {type(value).__name__}")
    return value


def _build_global_payload(ws) -> Dict[str, Any]:
    payload: Dict[str, Any] = {}
    for field, value, *_ in ws.iter_rows(min_row=2, values_only=True):
        if field is None:
            continue
        payload[str(field)] = _json_value(value)

    aliases = {
        "ppn_rate": payload.get("PPN Rate %"),
        "ppn_account": payload.get("PPN Account"),
        "dpp_account": payload.get("DPP Account"),
        "sales_tunai": payload.get("Sales Tunai Account"),
        "sales_qris_pending": payload.get("Sales QRIS Pending Account"),
        "sales_transfer": payload.get("Sales Transfer Account"),
        "selisih_plus": payload.get("Selisih Kas Plus Account"),
        "selisih_minus": payload.get("Selisih Kas Minus Account"),
        "mdr": payload.get("QRIS MDR % Default"),
        "qris_mdr": payload.get("QRIS MDR % Default"),
        "qris_enabled": payload.get("QRIS Enabled"),
        "tunai_enabled": payload.get("Tunai Enabled"),
        "transfer_enabled": payload.get("Transfer Enabled"),
        "currency": payload.get("Currency"),
        "timezone": payload.get("Timezone"),
        "package_id": payload.get("Package ID"),
    }
    payload.update({k: v for k, v in aliases.items() if v is not None})
    return payload


def _build_template_payloads(ws) -> Dict[str, Dict[str, Any]]:
    rows = _rows_as_dicts(ws)
    templates: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        template_id = str(row["template_id"]).strip().upper()
        fields = _parse_json_object(row.get("43_fields_json"), key=f"template:{template_id}")
        payload = {
            "template_id": template_id,
            "template_name": row.get("template_name"),
            "price": row.get("price_per_month"),
            "price_per_month": row.get("price_per_month"),
            "max_user": row.get("max_user"),
            "mdr": row.get("mdr_pos_%"),
            "mdr_pos": row.get("mdr_pos_%"),
            "mdr_akuntansi": row.get("mdr_akuntansi_%"),
            "features": {
                "pos": bool(row.get("features_pos")),
                "akuntansi": bool(row.get("features_akuntansi")),
                "api": bool(row.get("features_api")),
            },
            "tenants_count": row.get("tenants_count"),
            "description": row.get("description"),
            "fields": fields,
        }
        payload.update(fields)
        templates[template_id] = payload

    if set(templates) != EXPECTED_TEMPLATES:
        raise ValueError(f"Template IDs mismatch: got {sorted(templates)}")
    return templates


def _index_by(rows: Iterable[Dict[str, Any]], field: str) -> Dict[str, Dict[str, Any]]:
    out: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        value = row.get(field)
        if value is not None:
            out[str(value)] = row
    return out


def build_payloads(excel_path: Path) -> List[Dict[str, Any]]:
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    required = {
        "TEMPLATES", "GLOBAL", "TENANTS_MASTER", "DEVICES", "SERVICES", "APP_CONFIG_EXPORT"
    }
    missing = required.difference(wb.sheetnames)
    if missing:
        raise ValueError(f"Missing required sheets: {sorted(missing)}")

    global_payload = _build_global_payload(wb["GLOBAL"])
    template_payloads = _build_template_payloads(wb["TEMPLATES"])
    tenants = _rows_as_dicts(wb["TENANTS_MASTER"])
    devices = _rows_as_dicts(wb["DEVICES"])
    services = _rows_as_dicts(wb["SERVICES"])
    export_rows = _rows_as_dicts(wb["APP_CONFIG_EXPORT"])

    if len(tenants) != EXPECTED_TENANTS:
        raise ValueError(f"Expected {EXPECTED_TENANTS} tenants, got {len(tenants)}")
    if len(services) != EXPECTED_SERVICES:
        raise ValueError(f"Expected {EXPECTED_SERVICES} services, got {len(services)}")
    if len(devices) != EXPECTED_DEVICES:
        raise ValueError(f"Expected {EXPECTED_DEVICES} devices, got {len(devices)}")
    if len(export_rows) != EXPECTED_EXPORT_ROWS:
        raise ValueError(f"Expected {EXPECTED_EXPORT_ROWS} APP_CONFIG_EXPORT rows, got {len(export_rows)}")

    tenant_by_id = _index_by(tenants, "tenant_id")
    device_by_key = _index_by(devices, "app_config_key")
    service_by_key = _index_by(services, "app_config_key")

    payloads: List[Dict[str, Any]] = []
    seen = set()
    type_counts: Dict[str, int] = {}

    for row in export_rows:
        key = str(row.get("key") or "").strip()
        if not key:
            raise ValueError("APP_CONFIG_EXPORT contains an empty key")
        if key in seen:
            raise ValueError(f"Duplicate app_config key: {key}")
        seen.add(key)

        kind = str(row.get("type") or "").strip().lower()
        value = _parse_json_object(row.get("value_jsonb"), key=key)

        if key == "global":
            value = {**global_payload, **value}
        elif key.startswith("template:"):
            template_id = key.split(":", 1)[1].upper()
            value = {**template_payloads[template_id], **value}
            if "mdr_pos" in value:
                value["mdr"] = value["mdr_pos"]
        elif kind == "tenant":
            tenant_id = key.split(":", 1)[1]
            src = tenant_by_id.get(tenant_id, {})
            template_id = str(src.get("template_id") or value.get("template") or "").upper()
            value.update({
                "tenant_id": tenant_id,
                "template": template_id,
                "template_id": template_id,
                "nama_toko": src.get("nama_toko", value.get("nama_toko")),
                "email": src.get("email"),
                "max_user_override": src.get("max_user_override"),
                "mdr_pos_override": src.get("mdr_pos_override_%"),
                "features": {
                    "pos": bool(src.get("features_pos")),
                    "akuntansi": bool(src.get("features_akuntansi")),
                },
                "blocked": bool(src.get("blocked")),
                "alamat": src.get("alamat"),
                "device_id_default": src.get("device_id_default"),
                "licence_expiry": src.get("licence_expiry"),
                "status": src.get("status"),
            })
        elif kind == "service":
            src = service_by_key.get(key, {})
            mdr = src.get("mdr_%", value.get("mdr_%"))
            value.update({
                "tenant_id": src.get("tenant_id"),
                "template_id": src.get("template_id"),
                "service": src.get("service", value.get("service")),
                "enabled": bool(src.get("enabled")),
                "inherited_from_template": src.get("inherited_from_template"),
                "mdr_%": mdr,
                "mdr": mdr,
                "mdr_pos": mdr,
                "qris_mdr": mdr,
            })
        elif kind == "device":
            src = device_by_key.get(key, {})
            value.update({
                "device_id": src.get("device_id", value.get("device_id")),
                "tenant_id": src.get("tenant_id", value.get("tenant")),
                "template_id": src.get("template_id"),
                "nama_device": src.get("nama_device"),
                "is_active": bool(src.get("is_active")),
            })

        type_counts[kind] = type_counts.get(kind, 0) + 1
        payloads.append({"key": key, "value": value})

    expected_types = {"global": 1, "template": 3, "tenant": 100, "service": 140, "device": 108}
    if type_counts != expected_types:
        raise ValueError(f"APP_CONFIG_EXPORT type counts mismatch: {type_counts}")
    return payloads


def _chunks(items: List[Dict[str, Any]], size: int):
    for i in range(0, len(items), size):
        yield items[i:i + size]


def _canonical(v: Any) -> str:
    return json.dumps(v, sort_keys=True, separators=(",", ":"), ensure_ascii=False, default=str)


def _rest_headers(service_key: str, *, upsert: bool = False) -> Dict[str, str]:
    headers = {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
        "Accept": "application/json",
        "User-Agent": "berkahtigaputeri-final-pos-sync/1.0",
    }
    if upsert:
        headers["Prefer"] = "resolution=merge-duplicates,return=minimal"
    return headers


def _raise_rest_error(response: requests.Response, operation: str) -> None:
    if response.ok:
        return
    detail = response.text.strip().replace("\n", " ")[:1000]
    raise RuntimeError(
        f"Supabase REST {operation} failed: HTTP {response.status_code}. "
        f"{detail or 'No response body'}"
    )


def sync(payloads: List[Dict[str, Any]]) -> int:
    url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
    service_key = os.environ.get("SUPABASE_SERVICE_KEY", "").strip()
    if not url or not service_key:
        raise RuntimeError("SUPABASE_URL and SUPABASE_SERVICE_KEY are required")

    endpoint = f"{url}/rest/v1/app_config"
    read_response = requests.get(
        endpoint,
        headers=_rest_headers(service_key),
        params={"select": "key,value", "limit": "1000"},
        timeout=REQUEST_TIMEOUT,
    )
    _raise_rest_error(read_response, "read app_config")

    existing: Dict[str, Any] = {}
    for row in read_response.json():
        existing[str(row["key"])] = row.get("value")

    changed = [
        row for row in payloads
        if row["key"] not in existing or _canonical(existing[row["key"]]) != _canonical(row["value"])
    ]

    now = datetime.now(timezone.utc).isoformat()
    for batch in _chunks(changed, BATCH_SIZE):
        write_batch = [{**row, "updated_at": now} for row in batch]
        write_response = requests.post(
            endpoint,
            headers=_rest_headers(service_key, upsert=True),
            params={"on_conflict": "key"},
            json=write_batch,
            timeout=REQUEST_TIMEOUT,
        )
        _raise_rest_error(write_response, "upsert app_config")

    return len(changed)


def main() -> int:
    if not EXCEL_PATH.exists():
        print(f"ERROR: Excel not found: {EXCEL_PATH}", file=sys.stderr)
        return 2

    payloads = build_payloads(EXCEL_PATH)
    if os.environ.get("FINAL_POS_VALIDATE_ONLY") == "1":
        print(f"Validated {len(payloads)} keys - 100 tenants SaaS ready")
        return 0

    try:
        changed = sync(payloads)
    except (requests.RequestException, RuntimeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 3

    print(f"Synced {len(payloads)} keys - 100 tenants SaaS ready")
    print(f"Changed/upserted rows: {changed}; unchanged rows skipped: {len(payloads) - changed}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
